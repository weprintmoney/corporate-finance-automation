#!/usr/bin/env python3
"""Refresh the market-data snapshots in 03-corporate-finance/03-3-supplemental-data/.

Writes four JSON files. Each carries `as_of` (the date the publisher says the data describes),
`fetched` (the run date) and `source_url`:

  market-rates.json     From Damodaran's monthly ERP workbook: 10-year Treasury, US sovereign default
                        spread, dollar risk-free rate, and the implied equity risk premium over each.
                        From FRED: the latest 10-year Treasury, SOFR and Fed Funds closes.
  ratings-spreads.json  Interest coverage -> synthetic rating -> default spread (large, small, financial)
  country-risk.json     Per-country Moody's rating, default spread, equity and country risk premium
  industry-betas.json   US industry betas, levered and unlevered, cash-corrected

Usage:
  python3 scripts/refresh-market-data.py            fetch everything, rewrite the files, then --check
  python3 scripts/refresh-market-data.py --check    fetch nothing; exit 1 if a file is missing or stale
  python3 scripts/refresh-market-data.py --as-of YYYY-MM-DD --out DIR
                                                    the same four files as they stood on a past date,
                                                    written to DIR (never to the shared snapshot)

--as-of is for valuing a company at a past date, such as a delisted company on its last unaffected
trading day. It reads the ERP workbook for that month, the FRED closes on or before the date, and the
January editions of the ratings table, country risk and betas from Damodaran's archive. In --as-of mode
`fred_latest` means the latest close on or before the date.

Fetching needs `pip install xlrd openpyxl`; --check needs only the standard library. No API keys.

Every parser exits non-zero when the source layout changes, rather than writing an empty file. Damodaran
publishes the ERP workbook monthly, country risk in January and July, betas and the ratings table in
January. MAX_AGE_DAYS allows for that cadence plus some lag; a file past its window means a fetch broke
or a URL pattern changed, and /value-company refuses to use it.
"""
import calendar
import json
import os
import re
import sys
import tempfile
import urllib.request
from datetime import date, datetime, timedelta

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(REPO, "03-corporate-finance", "03-3-supplemental-data")

BASE = "https://pages.stern.nyu.edu/~adamodar"
RATINGS_XLS = BASE + "/pc/ratings.xls"
RATINGS_HTML = BASE + "/New_Home_Page/datafile/ratings.html"
CTRYPREM = BASE + "/pc/datasets/ctryprem.xlsx"                 # January update
CTRYPREM_JULY = BASE + "/pc/datasets/ctrypremJuly{yy}.xlsx"    # July update, separate file
BETAS = BASE + "/pc/datasets/betas.xls"
ERP_PATTERN = BASE + "/pc/implprem/ERP{month}{yy}.xlsx"
# Damodaran's archive names January editions by the prior year: betas24.xls is the 1/25 update.
ARCHIVE = {"ratings": BASE + "/pc/archives/ratings{yy}.xls", "ctryprem": BASE + "/pc/archives/ctryprem{yy}.xlsx",
           "betas": BASE + "/pc/archives/betas{yy}.xls"}
FRED_CSV = "https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}&cosd={start}&coed={end}"
FRED_SERIES = {"treasury_10y": "DGS10", "sofr": "SOFR", "fed_funds_effective": "DFF"}

MAX_AGE_DAYS = {
    "market-rates.json": 45,
    "country-risk.json": 200,
    "ratings-spreads.json": 400,
    "industry-betas.json": 400,
}


def fetch(url, dest):
    # Python's default user agent: FRED stalls on browser-like or custom ones
    with urllib.request.urlopen(url, timeout=60) as resp, open(dest, "wb") as fh:
        fh.write(resp.read())
    return dest


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def fail(msg):
    sys.exit("refresh-market-data: " + msg)


def value_right_of(rows, label):
    """First numeric cell to the right of the first cell whose text starts with `label`."""
    for row in rows:
        for i, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower().startswith(label.lower()):
                for nxt in row[i + 1:]:
                    if num(nxt) is not None:
                        return float(nxt)
    return None


# ---- market-rates.json ------------------------------------------------------------------------

def find_erp_workbook(tmp, today):
    """Damodaran names the file by month, inconsistently: ERPOct26, ERPJuly26, ERPSept25."""
    for back in range(4):
        y, m = today.year, today.month - back
        if m < 1:
            y, m = y - 1, m + 12
        full = calendar.month_name[m]
        for name in dict.fromkeys([full[:3], full[:4], full]):
            url = ERP_PATTERN.format(month=name, yy=str(y)[2:])
            dest = os.path.join(tmp, "erp.xlsx")
            try:
                fetch(url, dest)
            except Exception:
                continue
            with open(dest, "rb") as fh:
                if fh.read(2) == b"PK":          # an xlsx, not an error page
                    return url, dest, date(y, m, 1)
    fail("no ERP workbook for this month or the three before it ({})".format(ERP_PATTERN))


def fred_latest(series, tmp, today):
    start = (today - timedelta(days=21)).isoformat()
    path = fetch(FRED_CSV.format(series=series, start=start, end=today.isoformat()), os.path.join(tmp, series + ".csv"))
    rows = []
    for line in open(path):
        parts = line.strip().split(",")
        if len(parts) == 2 and parts[0][:1].isdigit() and num(parts[1]) is not None:
            rows.append((parts[0], float(parts[1]) / 100))
    if not rows:
        fail("FRED returned no observations for {}".format(series))
    return {"value": rows[-1][1], "date": rows[-1][0], "source_url": "https://fred.stlouisfed.org/series/" + series}


def market_rates(tmp, today):
    import openpyxl
    url, path, month = find_erp_workbook(tmp, today)
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["Impl premium calculator"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if month < US_DOWNGRADE_MONTH and value_right_of(rows, "Default spread for") is None:
        erp = pre_downgrade_rates(rows)
    else:
        erp = post_downgrade_rates(rows)
    if abs(erp["treasury_10y"] - erp["us_sovereign_default_spread"] - erp["dollar_riskfree_rate"]) > 0.0005:
        fail("ERP workbook is inconsistent: Treasury minus sovereign spread is not the dollar risk-free rate")
    return {
        "as_of": month.isoformat(),
        "fetched": date.today().isoformat(),
        "source_url": url,
        "note": ("damodaran_monthly values are as of the start of the month and pair with each other: "
                 "dollar_riskfree_rate with implied_erp_over_dollar_riskfree, or treasury_10y with "
                 "implied_erp_over_treasury; never mix. Pre-tax cost of debt = treasury_10y + the "
                 "company's default spread, because rating spreads are measured over Treasuries. "
                 "fred_latest is the most recent daily close, for checking how far rates moved since."),
        "damodaran_monthly": erp,
        "fred_latest": fred_all(tmp, today),
    }


# Workbooks before the June 2025 update predate Moody's downgrade of the US (2025-05-16): the US was
# rated Aaa, the sheet has a single risk-free rate and a single implied premium, and the spread is zero.
US_DOWNGRADE_MONTH = date(2025, 6, 1)


def pre_downgrade_rates(rows):
    tbond = value_right_of(rows, "Enter current long term risk free rate")
    implied = value_right_of(rows, "Implied Risk Premium in current level of Index")
    level = value_right_of(rows, "Enter current level of index")
    if None in (tbond, implied, level):
        fail("pre-downgrade ERP workbook layout changed; could not read the risk-free rate, premium or index level")
    return {"treasury_10y": tbond, "us_sovereign_default_spread": 0.0, "dollar_riskfree_rate": tbond,
            "implied_erp_over_dollar_riskfree": implied, "implied_erp_over_treasury": implied, "sp500_level": level}


def post_downgrade_rates(rows):
    erp = {
        "treasury_10y": value_right_of(rows, "US 10-year treasury rate ="),
        "us_sovereign_default_spread": value_right_of(rows, "Default spread for"),
        "dollar_riskfree_rate": value_right_of(rows, "Default-risk-adjusted US dollar riskfree"),
        "implied_erp_over_dollar_riskfree": value_right_of(rows, "Implied Equity Risk Premium (with default"),
        "implied_erp_over_treasury": value_right_of(rows, "Implied Equity Risk Premium (with US tre"),
        "sp500_level": value_right_of(rows, "Enter current level of index"),
    }
    missing = [k for k, v in erp.items() if v is None]
    if missing:
        fail("ERP workbook layout changed; could not read {}".format(", ".join(missing)))
    return erp


def fred_all(tmp, today):
    """A cross-check only, so a FRED outage is recorded rather than fatal."""
    out = {}
    for key, series in FRED_SERIES.items():
        try:
            out[key] = fred_latest(series, tmp, today)
        except Exception as exc:
            print("refresh-market-data: FRED {} unavailable ({})".format(series, exc))
            out[key] = {"value": None, "error": str(exc)}
    return out


# ---- ratings-spreads.json -----------------------------------------------------------------------

RATING_CELL = re.compile(r"^[A-Za-z0-9]+/[A-Z]+[+-]?$")      # "Aaa/AAA", "B3/B-"


def ratings_spreads(tmp, archive_year=None):
    import xlrd
    url = ARCHIVE["ratings"].format(yy=str(archive_year - 1)[2:]) if archive_year else RATINGS_XLS
    sheet = xlrd.open_workbook(fetch(url, os.path.join(tmp, "ratings.xls"))).sheet_by_index(0)
    blocks = {}                                               # (column, first row) -> rows
    open_block = {}                                           # column -> key of the block still growing
    for r in range(sheet.nrows):
        row = [c.value for c in sheet.row(r)]
        for c, cell in enumerate(row):
            if not (isinstance(cell, str) and RATING_CELL.match(cell.strip())) or c < 2 or c + 1 >= len(row):
                continue
            lo, spread = num(row[c - 2]), num(row[c + 1])
            if lo is None or spread is None:
                continue
            key = open_block.get(c)
            if key is None or blocks[key][-1]["row"] != r - 1:
                key = (c, r)
                blocks[key] = []
                open_block[c] = key
            blocks[key].append({"row": r, "min_coverage": lo, "rating": cell.strip(), "spread": spread})
    # Each 15-row table sits under a title such as "For smaller and riskier non-financial service firms".
    # Some editions add a utilities table; name every table by its title, then require the three we use.
    named = {}
    for (c, r), rows in blocks.items():
        if len(rows) != 15:
            continue
        title = " ".join(str(sheet.cell_value(rr, cc)).lower() for rr in range(max(0, r - 4), r)
                         for cc in range(max(0, c - 2), min(sheet.ncols, c + 2)))
        for key, words in [("small", ("smaller",)), ("utilities", ("infrastructure", "utilit")),
                           ("large", ("large",)), ("financial", ("financial",))]:
            if any(w in title for w in words):
                named.setdefault(key, (c, r))
                break
    if not all(k in named for k in ("large", "small", "financial")):
        fail("ratings.xls layout changed; found tables {} but need large, small and financial".format(sorted(named)))
    tables = {k: blocks[v] for k, v in named.items()}

    if archive_year:
        as_of = date(archive_year, 1, 1)                      # archived files are the January editions
    else:
        text = open(fetch(RATINGS_HTML, os.path.join(tmp, "ratings.html")), encoding="latin-1").read()
        m = re.search(r"as of\s+([A-Za-z]+)\s+(\d{4})", text)
        if not m:
            fail("ratings page no longer states 'as of <Month> <Year>'")
        as_of = datetime.strptime("{} {}".format(m.group(1), m.group(2)), "%B %Y").date()

    def clean(name):
        return [{k: v for k, v in row.items() if k != "row"} for row in tables[name]]

    return {
        "as_of": as_of.isoformat(),
        "fetched": date.today().isoformat(),
        "source_url": url,
        "note": ("Rows in ascending coverage; a firm gets the last row whose min_coverage its coverage "
                 "exceeds. large = market cap above about $5B. Spreads are over the 10-year Treasury, so "
                 "pre-tax cost of debt = Treasury rate + spread."),
        **{k: clean(k) for k in ("large", "small", "financial", "utilities") if k in tables},
    }


# ---- country-risk.json ---------------------------------------------------------------------------

def country_risk(tmp, archive_year=None, cutoff=None):
    """Damodaran posts the January update at ctryprem.xlsx and the July one at ctrypremJulyYY.xlsx; keep the newer."""
    if archive_year:
        return parse_country_risk(ARCHIVE["ctryprem"].format(yy=str(archive_year - 1)[2:]), tmp)
    year = date.today().year
    found = []
    for url in [CTRYPREM, CTRYPREM_JULY.format(yy=str(year)[2:]), CTRYPREM_JULY.format(yy=str(year - 1)[2:])]:
        try:
            found.append(parse_country_risk(url, tmp))
        except SystemExit as exc:
            print("refresh-market-data: skipped {} ({})".format(url, exc))
        except Exception as exc:
            print("refresh-market-data: could not fetch {} ({})".format(url, exc))
    if cutoff:                                                # --as-of: no edition published after the date
        found = [d for d in found if d["as_of"] <= cutoff.isoformat()]
    if not found:
        fail("no country risk workbook could be read")
    return max(found, key=lambda d: d["as_of"])


def parse_country_risk(url, tmp):
    import openpyxl
    path = fetch(url, os.path.join(tmp, "ctryprem.xlsx"))
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["ERPs by country"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    stamp = next((r[1] for r in rows[:5] if r and isinstance(r[0], str) and r[0].lower().startswith("date of update")), None)
    if not isinstance(stamp, datetime):
        fail("ctryprem.xlsx no longer states its update date on 'ERPs by country'")
    header = next((i for i, r in enumerate(rows) if r and r[0] == "Country"), None)
    if header is None:
        fail("ctryprem.xlsx has no 'Country' header row on 'ERPs by country'")
    countries = {}
    for row in rows[header + 1:]:
        if not row or not row[0]:
            break
        if num(row[4]) is not None:
            countries[str(row[0]).strip()] = {
                "region": row[1], "moodys_rating": row[2], "default_spread": num(row[3]),
                "erp": num(row[4]), "country_risk_premium": num(row[5]),
            }
    if len(countries) < 100:
        fail("ctryprem.xlsx parsed only {} countries".format(len(countries)))
    return {
        "as_of": stamp.date().isoformat(),
        "fetched": date.today().isoformat(),
        "source_url": url,
        "note": "Rating-based default spreads and premiums. Damodaran updates in January and July.",
        "mature_market_erp": value_right_of(rows[:6], "Enter the current risk premium for a mature"),
        "us_erp": value_right_of(rows[:6], "Enter the current risk premium for the US"),
        "countries": countries,
    }


# ---- industry-betas.json -------------------------------------------------------------------------

def industry_betas(tmp, archive_year=None):
    import xlrd
    url = ARCHIVE["betas"].format(yy=str(archive_year - 1)[2:]) if archive_year else BETAS
    book = xlrd.open_workbook(fetch(url, os.path.join(tmp, "betas.xls")))
    sheet = book.sheet_by_name("Industry Averages")
    rows = [[c.value for c in sheet.row(r)] for r in range(sheet.nrows)]
    if not (rows and str(rows[0][0]).lower().startswith("date updated") and num(rows[0][1])):
        fail("betas.xls no longer states 'Date updated' in its first row")
    as_of = xlrd.xldate.xldate_as_datetime(num(rows[0][1]), book.datemode).date()
    header = next((i for i, r in enumerate(rows) if r and r[0] == "Industry Name"), None)
    if header is None:
        fail("betas.xls has no 'Industry Name' header row")
    industries = {}
    for row in rows[header + 1:]:
        if not row[0] or num(row[1]) is None:
            break
        industries[str(row[0]).strip()] = {
            "firms": int(num(row[1])), "beta": num(row[2]), "de_ratio": num(row[3]),
            "effective_tax_rate": num(row[4]), "unlevered_beta": num(row[5]),
            "cash_to_firm_value": num(row[6]), "unlevered_beta_cash_corrected": num(row[7]),
        }
    if len(industries) < 50:
        fail("betas.xls parsed only {} industries".format(len(industries)))
    return {
        "as_of": as_of.isoformat(),
        "fetched": date.today().isoformat(),
        "source_url": url,
        "note": "US companies. Damodaran refreshes in January.",
        "industries": industries,
    }


# ---- staleness check -----------------------------------------------------------------------------

def check(today=None):
    today = today or date.today()
    bad = []
    for name, limit in MAX_AGE_DAYS.items():
        try:
            with open(os.path.join(OUT_DIR, name)) as fh:
                as_of = datetime.strptime(json.load(fh)["as_of"], "%Y-%m-%d").date()
        except Exception:
            print("{:<22} missing or has no as_of".format(name))
            bad.append(name)
            continue
        age = (today - as_of).days
        print("{:<22} as_of {}  {:>4} days old  (limit {})  {}".format(name, as_of, age, limit, "ok" if age <= limit else "STALE"))
        if age > limit:
            bad.append(name)
    if bad:
        print("Stale or missing: {}. Run python3 scripts/refresh-market-data.py.".format(", ".join(bad)))
        return 1
    return 0


def main(argv):
    if "--check" in argv:
        sys.exit(check())
    as_of, out_dir = None, OUT_DIR
    if "--as-of" in argv:
        try:
            as_of = datetime.strptime(argv[argv.index("--as-of") + 1], "%Y-%m-%d").date()
            out_dir = argv[argv.index("--out") + 1]
        except (IndexError, ValueError):
            sys.exit("usage: refresh-market-data.py --as-of YYYY-MM-DD --out DIR")
        os.makedirs(out_dir, exist_ok=True)
    # The current year's January editions are the live files; earlier years come from the archive.
    archive_year = as_of.year if as_of and as_of.year < date.today().year else None
    with tempfile.TemporaryDirectory() as tmp:
        outputs = {
            "market-rates.json": market_rates(tmp, as_of or date.today()),
            "ratings-spreads.json": ratings_spreads(tmp, archive_year),
            "country-risk.json": country_risk(tmp, archive_year, as_of),
            "industry-betas.json": industry_betas(tmp, archive_year),
        }
    for name, payload in outputs.items():           # write only after every source parsed
        if as_of:
            payload["requested_as_of"] = as_of.isoformat()
        with open(os.path.join(out_dir, name), "w") as fh:
            json.dump(payload, fh, indent=2)
            fh.write("\n")
        print("wrote {} (as_of {})".format(os.path.join(out_dir, name), payload["as_of"]))
    if not as_of:
        sys.exit(check())


if __name__ == "__main__":
    main(sys.argv[1:])
